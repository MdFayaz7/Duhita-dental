"""Read a Practo patient export (CSV, XLSX or XLS) into clean patient rows.

Practo's column names vary a little between export screens and versions, so headers are matched
by alias; the admin can still override the mapping in the import screen before anything is saved.
"""
import csv
import datetime as dt
import io
import re
from collections import Counter

PHONE_RE = re.compile(r"^[6-9]\d{9}$")

# field -> header aliases, compared with everything but letters and digits removed
ALIASES: dict[str, list[str]] = {
    "practo_id": ["patientnumber", "patientno", "patientid", "practoid", "uhid", "patientcode", "mrn", "casenumber"],
    "name": ["patientname", "name", "fullname", "patient"],
    "phone": ["mobilenumber", "mobile", "mobileno", "phone", "phonenumber", "primarymobile", "primarynumber",
              "contactnumber", "contact"],
    "alt_phone": ["secondarymobile", "secondarymobilenumber", "alternatemobile", "alternatenumber",
                  "alternatephone", "secondarynumber", "landline", "othernumber", "contactnumber2"],
    "email": ["emailaddress", "email", "emailid", "mail"],
    "sex": ["gender", "sex"],
    "dob": ["dateofbirth", "dob", "birthdate", "birthday"],
    "age": ["age", "ageyears"],
    "address": ["address", "streetaddress", "address1", "addressline1"],
    "locality": ["locality", "area", "address2", "addressline2"],
    "city": ["city", "town"],
    "pincode": ["pincode", "pin", "zipcode", "zip", "postalcode"],
    "blood_group": ["bloodgroup", "blood"],
    "referral": ["referredby", "referral", "reference", "source", "howdidyouhearaboutus"],
    "medical_history": ["medicalhistory", "medicalconditions", "conditions", "history"],
    "allergies": ["allergies", "allergy", "drugallergies"],
    "notes": ["patientnotes", "notes", "remarks", "comments"],
    "groups": ["groups", "group", "patientgroups", "tags"],
    "registered_on": ["registeredon", "registrationdate", "dateofregistration", "createdon", "createdat",
                      "firstvisit", "firstvisitdate", "addedon"],
}
FIELDS = list(ALIASES)
REQUIRED = ("name",)

DATE_FORMATS = ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y", "%d %b %Y", "%d-%b-%Y", "%d %B %Y",
                "%b %d, %Y", "%d/%m/%y", "%d-%m-%y", "%Y/%m/%d")
BLOOD_GROUPS = {"A+", "A-", "B+", "B-", "AB+", "AB-", "O+", "O-"}


class ImportFileError(ValueError):
    pass


def _key(header: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(header).lower())


# ---------------------------------------------------------------- reading

def read_table(filename: str, data: bytes) -> tuple[list[str], list[list]]:
    """Return (headers, rows) from a CSV / XLSX / XLS file. Blank leading rows are skipped."""
    name = (filename or "").lower()
    if name.endswith(".xlsx") or data[:2] == b"PK":
        rows = _read_xlsx(data)
    elif name.endswith(".xls") or data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        rows = _read_xls(data)
    else:
        rows = _read_csv(data)

    rows = [r for r in rows if any(str(c).strip() for c in r if c is not None)]
    if not rows:
        raise ImportFileError("The file is empty.")
    # Practo sometimes puts a title line above the real headers: use the first row that names a known field.
    for i, row in enumerate(rows[:10]):
        if sum(1 for c in row if c is not None and _match(str(c))) >= 2:
            headers = [str(c).strip() if c is not None else "" for c in row]
            return headers, rows[i + 1:]
    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    return headers, rows[1:]


def _read_csv(data: bytes) -> list[list]:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text), dialect))


def _read_xlsx(data: bytes) -> list[list]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFileError(f"Could not open the Excel file: {exc}")
    sheet = wb.worksheets[0]
    rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    wb.close()
    return rows


def _read_xls(data: bytes) -> list[list]:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as exc:
        raise ImportFileError(f"Could not open the Excel file: {exc}")
    sheet = book.sheet_by_index(0)
    out = []
    for r in range(sheet.nrows):
        row = []
        for cell in sheet.row(r):
            if cell.ctype == xlrd.XL_CELL_DATE:
                row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
            elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                row.append(int(cell.value))
            else:
                row.append(cell.value)
        out.append(row)
    return out


# ---------------------------------------------------------------- mapping

def _match(header: str) -> str | None:
    key = _key(header)
    for field, aliases in ALIASES.items():
        if key in aliases:
            return field
    return None


def detect_mapping(headers: list[str]) -> dict[str, str]:
    """field -> header, for every header we recognise (first match wins)."""
    mapping: dict[str, str] = {}
    for header in headers:
        field = _match(header)
        if field and field not in mapping:
            mapping[field] = header
    return mapping


# ---------------------------------------------------------------- cleaning

def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def phones_in(value) -> list[str]:
    """Every valid Indian mobile number in a cell ('+91 98765 43210 / 09123456789' -> two numbers)."""
    text = _text(value)
    found = []
    for chunk in re.split(r"[,/;|&]| or ", text):
        digits = re.sub(r"\D", "", chunk)
        if len(digits) == 12 and digits.startswith("91"):
            digits = digits[2:]
        elif len(digits) == 11 and digits.startswith("0"):
            digits = digits[1:]
        if PHONE_RE.match(digits) and digits not in found:
            found.append(digits)
    return found


def parse_date(value) -> dt.date | None:
    if value in (None, ""):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = _text(value)
    text = re.split(r"[ T](?=\d{1,2}:\d{2})", text)[0]  # drop a time part
    for fmt in DATE_FORMATS:
        try:
            day = dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
        return day if 1900 <= day.year <= dt.date.today().year + 1 else None
    return None


def _sex(value) -> str | None:
    text = _text(value).lower()
    if text in ("m", "male", "man", "boy"):
        return "Male"
    if text in ("f", "female", "woman", "girl"):
        return "Female"
    if text in ("o", "other", "others", "transgender"):
        return "Other"
    return None


def _age(value, dob: dt.date | None) -> int | None:
    if dob:
        today = dt.date.today()
        return max(0, min(120, today.year - dob.year - ((today.month, today.day) < (dob.month, dob.day))))
    match = re.match(r"\d{1,3}", _text(value))
    if match and int(match.group()) <= 120:
        return int(match.group())
    return None


def _name(value) -> str:
    text = _text(value)
    # Practo exports are often ALL CAPS; mixed-case names are left as typed.
    return text.title() if text.isupper() else text


def clean_row(raw: dict[str, object]) -> dict:
    """One export row (field -> cell) -> patient fields, plus `_issues` for the preview."""
    issues = []
    name = _name(raw.get("name"))
    if len(name) < 2:
        issues.append("no name")

    numbers = phones_in(raw.get("phone"))
    for extra_number in phones_in(raw.get("alt_phone")):
        if extra_number not in numbers:
            numbers.append(extra_number)
    if not numbers:
        issues.append("no valid mobile number")

    dob = parse_date(raw.get("dob"))
    address = ", ".join(p for p in (_text(raw.get(k)) for k in ("address", "locality", "city", "pincode")) if p)
    blood = _text(raw.get("blood_group")).upper().replace(" ", "").replace("VE", "")
    history = [h.strip() for h in re.split(r"[,;\n]", _text(raw.get("medical_history"))) if h.strip()]
    registered = parse_date(raw.get("registered_on"))
    extra = {k: v for k, v in (("practo_notes", _text(raw.get("notes"))), ("practo_groups", _text(raw.get("groups"))))
             if v}

    return {
        "practo_id": _text(raw.get("practo_id")) or None,
        "name": name,
        "phone": numbers[0] if numbers else "",
        "alt_phone": numbers[1] if len(numbers) > 1 else None,
        "email": (_text(raw.get("email")).lower() or None) if "@" in _text(raw.get("email")) else None,
        "sex": _sex(raw.get("sex")),
        "dob": dob,
        "age": _age(raw.get("age"), dob),
        "address": address or None,
        "blood_group": blood if blood in BLOOD_GROUPS else None,
        "referral": _text(raw.get("referral")) or None,
        "conditions": history,
        "allergies": _text(raw.get("allergies")) or None,
        "registered_on": registered,
        "extra": extra,
        "_issues": issues,
    }


def clean_rows(headers: list[str], rows: list[list], mapping: dict[str, str]) -> list[dict]:
    index = {h: i for i, h in enumerate(headers)}
    columns = {field: index[h] for field, h in mapping.items() if field in ALIASES and h in index}
    out = []
    for row in rows:
        raw = {field: (row[i] if i < len(row) else None) for field, i in columns.items()}
        out.append(clean_row(raw))
    return out


def family_phones(rows: list[dict]) -> Counter:
    """Mobile numbers shared by more than one patient in the file."""
    counts = Counter(r["phone"] for r in rows if r["phone"])
    return Counter({p: n for p, n in counts.items() if n > 1})

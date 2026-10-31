"""Excel export of the patient lists (clinic patients or the Practo archive) with the columns the admin picks."""
import datetime as dt
import io

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import func, select, text

from ..db import SessionDep
from ..security import current_admin
from ..tables import Invoice, Patient, PractoPatient

router = APIRouter(prefix="/api/export", tags=["export"], dependencies=[Depends(current_admin)])

_join = lambda v: ", ".join(v) if v else None  # noqa: E731

# key -> (heading, value from the row); clinic rows also carry billed / paid totals
CLINIC = {
    "patient_id": ("Patient ID", lambda p, b: p.patient_id),
    "name": ("Name", lambda p, b: p.name),
    "phone": ("Mobile", lambda p, b: p.phone),
    "alt_phone": ("Alternate mobile", lambda p, b: p.alt_phone),
    "age": ("Age", lambda p, b: p.age),
    "sex": ("Gender", lambda p, b: p.sex),
    "dob": ("Date of birth", lambda p, b: p.dob),
    "email": ("Email", lambda p, b: p.email),
    "address": ("Address", lambda p, b: p.address),
    "blood_group": ("Blood group", lambda p, b: p.blood_group),
    "conditions": ("Medical conditions", lambda p, b: _join(p.conditions)),
    "allergies": ("Allergies", lambda p, b: p.allergies),
    "medications": ("Current medicines", lambda p, b: p.medications),
    "complaint": ("Complaint", lambda p, b: p.complaint),
    "referral": ("Referred by", lambda p, b: " — ".join(filter(None, (p.referral, p.referral_name))) or None),
    "profession": ("Profession", lambda p, b: p.profession),
    "emergency": ("Emergency contact", lambda p, b: " ".join(filter(None, (p.emergency_contact, p.emergency_phone))) or None),
    "source": ("Registered via", lambda p, b: p.source),
    "app_user": ("App account", lambda p, b: "Yes" if p.app_user else "No"),
    "practo_id": ("Practo number", lambda p, b: p.practo_id),
    "created_at": ("Registered on", lambda p, b: p.created_at.date() if p.created_at else None),
    "billed": ("Total billed (₹)", lambda p, b: b[0]),
    "paid": ("Total paid (₹)", lambda p, b: b[1]),
    "due": ("Due (₹)", lambda p, b: b[0] - b[1]),
}

PRACTO = {
    "number": ("Patient Number", lambda p: p.number if not p.number.startswith("NA-") else None),
    "name": ("Patient Name", lambda p: p.name),
    "mobile": ("Mobile Number", lambda p: p.mobile),
    "gender": ("Gender", lambda p: p.gender),
    "age": ("Age", lambda p: p.age),
    "dob": ("Date of Birth", lambda p: p.dob),
    "email": ("Email Address", lambda p: p.email),
    "address": ("Address", lambda p: p.address),
    "records": ("Records", lambda p: p.records),
    "last_visit": ("Last visit", lambda p: p.last_visit),
}


async def _practo_extra(session) -> list[str]:
    """The other columns of Practo's own patient export, kept in `data` (e.g. Blood Group, Medical History)."""
    rows = await session.execute(text("SELECT DISTINCT jsonb_object_keys(data) FROM practo_patients ORDER BY 1"))
    return [r[0] for r in rows]


@router.get("/columns")
async def columns(session: SessionDep, kind: str = Query("clinic", pattern="^(clinic|practo)$")):
    if kind == "clinic":
        return {"columns": [{"key": k, "label": v[0]} for k, v in CLINIC.items()]}
    extra = await _practo_extra(session)
    return {"columns": [{"key": k, "label": v[0]} for k, v in PRACTO.items()] + [{"key": f"data:{k}", "label": k} for k in extra]}


@router.get("/patients")
async def export_patients(session: SessionDep, kind: str = Query("clinic", pattern="^(clinic|practo)$"),
                          columns: str = Query(..., description="comma-separated column keys, in order")):
    keys = [k for k in columns.split(",") if k]
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Patients")

    if kind == "clinic":
        if not keys or any(k not in CLINIC for k in keys):
            raise HTTPException(422, "Unknown column.")
        totals = {pid: (t or 0, p or 0) for pid, t, p in (await session.execute(
            select(Invoice.patient_id, func.sum(Invoice.total), func.sum(Invoice.paid))
            .where(Invoice.status == "active").group_by(Invoice.patient_id))).all()}
        ws.append([CLINIC[k][0] for k in keys])
        for p in await session.scalars(select(Patient).order_by(Patient.created_at.desc(), Patient.id.desc())):
            b = totals.get(p.patient_id, (0, 0))
            ws.append([CLINIC[k][1](p, b) for k in keys])
    else:
        extra = set(await _practo_extra(session))
        if not keys or any(k not in PRACTO and k.removeprefix("data:") not in extra for k in keys):
            raise HTTPException(422, "Unknown column.")
        label = lambda k: PRACTO[k][0] if k in PRACTO else k.removeprefix("data:")  # noqa: E731
        value = lambda p, k: PRACTO[k][1](p) if k in PRACTO else (p.data or {}).get(k.removeprefix("data:"))  # noqa: E731
        ws.append([label(k) for k in keys])
        for p in await session.scalars(select(PractoPatient).order_by(PractoPatient.name)):
            ws.append([value(p, k) for k in keys])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"{'practo-' if kind == 'practo' else ''}patients-{dt.date.today():%Y-%m-%d}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})
